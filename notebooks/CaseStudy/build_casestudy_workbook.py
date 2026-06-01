#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Build an Excel companion workbook for the TIPS-Treasury arbitrage Case Study
from Lecture 04 (04_TIPSTreasTrade.jl, lines 1349-2298).

Trade date: 5 October 2006.
  Long  TIPS 912828EA  (1.875%, 7/15/2015)
  + pay-fixed inflation swaps
  + Treasury STRIPS
  Short Treasury note 912828EE (4.250%, 8/15/2015)

The workbook reproduces the worked example as LIVE Excel formulas across 5 sheets
(Securities / MarketData / Step1_Prices / Step2_CashFlows / Result) and embeds the
6 Bloomberg screenshots that appear in the Case Study section.

Toolchain: Python + openpyxl + Pillow (read-only Pillow use, just to size images).
Run:  python build_casestudy_workbook.py
Output: TIPS_Treasury_CaseStudy.xlsx (next to this script)

The script also recomputes the example in pure Python and asserts the results match
the notebook's published checkpoints before saving (openpyxl stores formulas but does
not evaluate them, so this is the numeric gate).
"""

import os
import sys
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

# --------------------------------------------------------------------------------------
# Paths
# --------------------------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.normpath(os.path.join(HERE, "..", "Assets"))
# Output path; override with a command-line argument (used for diagnostic temp builds).
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "TIPS_Treasury_CaseStudy.xlsx")

SCREENSHOTS = [
    ("TIPS 912828EA - Description (Bloomberg DES)",        "TIPS_912828EA_Des.png"),
    ("TIPS 912828EA - Price quotes (Bloomberg)",           "TIPS_912828EA_PxQuotes.png"),
    ("TIPS 912828EA - Price history (Bloomberg GP)",       "TIPS_912828EA_PxGraph.png"),
    ("Treasury note 912828EE - Description (Bloomberg DES)","Treasury_912828EE_Des.png"),
    ("Treasury note 912828EE - Price quotes (Bloomberg)",  "Treasury_912828EE_PxQuotes.png"),
    ("Treasury note 912828EE - Price history (Bloomberg GP)","Treasury_912828EE_PxGraph_2.png"),
]

# --------------------------------------------------------------------------------------
# Case Study input data  (verbatim from 04_TIPSTreasTrade.jl)
# --------------------------------------------------------------------------------------
TRADE_DATE = "5 October 2006"

# Coupon cash-flow dates of the TIPS  (l.2026)
CF_DATES = [
    dt.date(2007, 1, 15), dt.date(2007, 7, 15), dt.date(2008, 1, 15), dt.date(2008, 7, 15),
    dt.date(2009, 1, 15), dt.date(2009, 7, 15), dt.date(2010, 1, 15), dt.date(2010, 7, 15),
    dt.date(2011, 1, 15), dt.date(2011, 7, 15), dt.date(2012, 1, 15), dt.date(2012, 7, 15),
    dt.date(2013, 1, 15), dt.date(2013, 7, 15), dt.date(2014, 1, 15), dt.date(2014, 7, 15),
    dt.date(2015, 1, 15), dt.date(2015, 7, 15),
]
# Interpolated, seasonally-adjusted inflation swap rates at the TIPS coupon dates (l.2027)
CF_SWAPS = [0.014123, 0.017091, 0.019628, 0.020925, 0.021960, 0.022847, 0.023544, 0.024078,
            0.024534, 0.024914, 0.025180, 0.025364, 0.025533, 0.025768, 0.025953, 0.026201,
            0.026308, 0.026458]
# Swap tenors as a fraction of a year (l.2028)
CF_TENORS = [0.279500, 0.775300, 1.279500, 1.778100, 2.282200, 2.778100, 3.282200, 3.778100,
             4.282200, 4.778100, 5.282200, 5.780800, 6.284900, 6.780800, 7.284900, 7.780800,
             8.284900, 8.780800]
# Interpolated Treasury STRIPS prices at the TIPS coupon dates (l.2205)
STRIPS_PX = [98.6650, 96.3040, 94.1340, 92.0890, 90.1020, 88.1370, 86.2470, 84.4300,
             82.6180, 80.9500, 78.6780, 76.8630, 75.0760, 73.3200, 71.5890, 69.9060,
             68.1530, 66.5100]
N = len(CF_DATES)  # 18

# Raw Bloomberg inflation-swap curve on 10/5/2006 (reference only, l.1447-1464)
RAW_SWAP_TENORS = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 12, 15, 20, 25, 30]
RAW_SWAP_RATES = [0.01850, 0.02133, 0.02313, 0.02425, 0.02503, 0.02543, 0.02580, 0.02623,
                  0.02648, 0.02683, 0.02733, 0.02798, 0.02960, 0.03070, 0.03170]

# Step-1 price inputs
TSY_COUPON, TSY_DSL, TSY_DCP = 0.0425, 51, 184    # coupon, days since last, days in period
TIPS_COUPON, TIPS_DSL, TIPS_DCP = 0.01875, 82, 184

# Settlement / coupon dates
SETTLE = dt.date(2006, 10, 5)
TSY_LAST_CPN, TSY_NEXT_CPN = dt.date(2006, 8, 15), dt.date(2007, 2, 15)
TIPS_LAST_CPN, TIPS_NEXT_CPN = dt.date(2006, 7, 15), dt.date(2007, 1, 15)

# Accrued-interest day-count detail (Actual/Actual), verbatim from the notebook
# (Treasury l.1637-1668, TIPS l.1730-1761). Each (label, days) row is a live input that
# sums to the numerator (days since last coupon) / denominator (days in coupon period).
TSY_DAYS_SINCE = [("August 2006  - incl. last coupon date (16+1)", 17),
                  ("September 2006", 30),
                  ("October 2006  - excl. settlement date", 4)]
TSY_DAYS_PERIOD = [("August 2006  (31-15)", 16), ("September 2006", 30), ("October 2006", 31),
                   ("November 2006", 30), ("December 2006", 31), ("January 2007", 31),
                   ("February 2007  (to the 15th)", 15)]
TIPS_DAYS_SINCE = [("July 2006  - incl. last coupon date (16+1)", 17),
                   ("August 2006", 31), ("September 2006", 30),
                   ("October 2006  - excl. settlement date", 4)]
TIPS_DAYS_PERIOD = [("July 2006  (31-15)", 16), ("August 2006", 31), ("September 2006", 30),
                    ("October 2006", 31), ("November 2006", 30), ("December 2006", 31),
                    ("January 2007  (to the 15th)", 15)]

# Maturity-adjustment figures (notebook "Details", l.2293-2297)
TIPS_YIELD = 0.0497
TSY_REPRICED_AT_TIPS_YIELD = 95.4583

# --------------------------------------------------------------------------------------
# Styling helpers
# --------------------------------------------------------------------------------------
NAVY = "1F3864"
LIGHTBLUE = "D9E1F2"
GREY = "F2F2F2"
GOLD = "FFF2CC"

TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
H2_FONT = Font(bold=True, size=12, color=NAVY)
HDR_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
ITAL = Font(italic=True, size=9, color="595959")

TITLE_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FILL = PatternFill("solid", fgColor="4472C4")
INPUT_FILL = PatternFill("solid", fgColor=GOLD)        # yellow = hard input
CALC_FILL = PatternFill("solid", fgColor=LIGHTBLUE)    # blue  = formula
RESULT_FILL = PatternFill("solid", fgColor="C6E0B4")   # green = headline result

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_PRICE = "0.0000"
FMT_ACCR = "0.000000"
FMT_RATE = "0.000000"
FMT_TENOR = "0.0000"
FMT_PCT = "0.000%"
FMT_DATE = "mm/dd/yyyy"


def title(ws, text, ncols=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, text)
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24


def h2(ws, row, text):
    c = ws.cell(row, 1, text)
    c.font = H2_FONT
    return c


def note(ws, row, text, col=1, span=8):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    c = ws.cell(row, col, text)
    c.font = ITAL
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    lines = max(1, (len(text) + 69) // 70)
    ws.row_dimensions[row].height = 15 * lines + 3


def hdr(ws, row, col, text):
    c = ws.cell(row, col, text)
    c.font = HDR_FONT
    c.fill = HDR_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    return c


def fix_empty_formula_values(xlsx_path):
    """openpyxl serializes every formula cell as <f>...</f><v/> with an EMPTY cached value.
    Excel treats that empty <v/> as corrupt content and shows a 'repair' prompt on open.
    Replace each empty <v/> with <v>0</v> (a placeholder, exactly as xlsxwriter does); combined
    with fullCalcOnLoad=True, Excel recomputes the real values on open and the file opens cleanly
    with no warning. Only worksheet XML is touched; all other parts are copied byte-for-byte."""
    import zipfile
    tmp = xlsx_path + ".fixtmp"
    with zipfile.ZipFile(xlsx_path, "r") as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                text = data.decode("utf-8").replace("<v/>", "<v>0</v>").replace("<v />", "<v>0</v>")
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, xlsx_path)


# --------------------------------------------------------------------------------------
# Build workbook
# --------------------------------------------------------------------------------------
wb = Workbook()

# ===== Sheet 1: Securities ============================================================
ws = wb.active
ws.title = "Securities"
title(ws, "TIPS-Treasury Arbitrage  -  Case Study   (trade date %s)" % TRADE_DATE, ncols=6)
ws.column_dimensions["A"].width = 26
for col in "BCDEF":
    ws.column_dimensions[col].width = 17

h2(ws, 3, "The two securities")
facts = [
    ("",                    "TIPS",       "Treasury note"),
    ("CUSIP",               "912828EA",   "912828EE"),
    ("Maturity",            "7/15/2015",  "8/15/2015"),
    ("Issue date",          "7/15/2005",  "8/15/2005"),
    ("Coupon",              "1.875%",     "4.250%"),
    ("First coupon date",   "1/15/2006",  "2/15/2006"),
    ("Second coupon date",  "7/15/2006",  "8/15/2006"),
    ("Reference CPI on issue date", "194.50968", "-"),
    ("Price quote",         "96-20",      "97-15+"),
]
r0 = 4
for i, (lab, tips, tsy) in enumerate(facts):
    r = r0 + i
    a = ws.cell(r, 1, lab); a.font = BOLD if i else BOLD
    b = ws.cell(r, 2, tips)
    c = ws.cell(r, 3, tsy)
    for cc in (a, b, c):
        cc.border = BORDER
    if i == 0:
        for cc in (b, c):
            cc.font = HDR_FONT
            cc.fill = HDR_FILL
            cc.alignment = Alignment(horizontal="center")
    else:
        for cc in (b, c):
            cc.alignment = Alignment(horizontal="center")
        a.fill = PatternFill("solid", fgColor=GREY)

note(ws, r0 + len(facts) + 1,
     "Bloomberg screenshots below are taken on the trade date. "
     "The pricing math is on the Step1_Prices / Step2_CashFlows / Result tabs.",
     span=6)

# Embed the 6 screenshots, stacked, scaled to a target display width.
img_row = r0 + len(facts) + 3
TARGET_W = 680  # px
for label, fname in SCREENSHOTS:
    path = os.path.join(ASSETS, fname)
    if not os.path.exists(path):
        ws.cell(img_row, 1, "[missing image: %s]" % fname).font = ITAL
        img_row += 2
        continue
    lab = ws.cell(img_row, 1, label)
    lab.font = BOLD
    with PILImage.open(path) as im:
        w, h = im.size
    disp_w = TARGET_W
    disp_h = int(round(TARGET_W * h / w))
    xfile = XLImage(path)
    xfile.width = disp_w
    xfile.height = disp_h
    anchor = "A%d" % (img_row + 1)
    ws.add_image(xfile, anchor)
    # advance roughly past the image (≈18 px per row) plus a gap
    img_row += int(disp_h / 18) + 3

# ===== Sheet 2: MarketData ============================================================
md = wb.create_sheet("MarketData")
title(md, "Market data on the trade date (%s)" % TRADE_DATE, ncols=5)
md.column_dimensions["A"].width = 14
for col in "BCDE":
    md.column_dimensions[col].width = 18

# (a) Raw Bloomberg swap curve - reference
h2(md, 3, "(a) Bloomberg inflation-swap curve  -  reference")
hdr(md, 4, 1, "Tenor (yrs)")
hdr(md, 4, 2, "Swap rate")
for i, (t, rr) in enumerate(zip(RAW_SWAP_TENORS, RAW_SWAP_RATES)):
    r = 5 + i
    a = md.cell(r, 1, t); a.alignment = Alignment(horizontal="center"); a.border = BORDER
    b = md.cell(r, 2, rr); b.number_format = FMT_RATE; b.fill = INPUT_FILL; b.border = BORDER

# (b) Interpolated inputs used by the calculation
MD_HDR_ROW = 4 + len(RAW_SWAP_TENORS) + 2  # leave a gap
h2(md, MD_HDR_ROW - 1, "(b) Inputs at the 18 TIPS coupon dates  -  used by Step2")
hdr(md, MD_HDR_ROW, 1, "Date")
hdr(md, MD_HDR_ROW, 2, "Inflation swap rate")
hdr(md, MD_HDR_ROW, 3, "Swap tenor (yrs)")
hdr(md, MD_HDR_ROW, 4, "STRIPS price")
MD_DATA0 = MD_HDR_ROW + 1
for i in range(N):
    r = MD_DATA0 + i
    a = md.cell(r, 1, CF_DATES[i]); a.number_format = FMT_DATE; a.border = BORDER
    b = md.cell(r, 2, CF_SWAPS[i]); b.number_format = FMT_RATE; b.fill = INPUT_FILL; b.border = BORDER
    c = md.cell(r, 3, CF_TENORS[i]); c.number_format = FMT_TENOR; c.fill = INPUT_FILL; c.border = BORDER
    d = md.cell(r, 4, STRIPS_PX[i]); d.number_format = FMT_PRICE; d.fill = INPUT_FILL; d.border = BORDER
MD_DATA_LAST = MD_DATA0 + N - 1
note(md, MD_DATA_LAST + 2,
     "Swap rates/tenors and STRIPS prices are interpolated to the TIPS coupon dates and "
     "seasonally adjusted; see Fleckenstein, Longstaff & Lustig (2014). Yellow = input.",
     span=4)
note(md, MD_DATA_LAST + 4,
     "Swap tenor day-count example (first row, 0.2795 yr): (31-5) days in Oct + 30 Nov + 31 Dec "
     "+ 15 Jan = 102 days, and 102 / 365 = 0.2795. The other tenors are computed the same way.",
     span=4)

# ===== Sheet 3: Step1_Prices ==========================================================
s1 = wb.create_sheet("Step1_Prices")
title(s1, "Step 1  -  Full (dirty) prices = quoted price + accrued interest", ncols=4)
s1.column_dimensions["A"].width = 46
s1.column_dimensions["B"].width = 15
s1.column_dimensions["C"].width = 14


def price_block(ws, r, header, quote_label, quote_formula, quote_note, coupon,
                settle, last_cpn, next_cpn, days_since_rows, days_period_rows):
    """Lay out one security's full-price calculation including the day-count detail.
    The monthly day counts are live inputs that SUM to the accrued numerator/denominator.
    Returns (full_price_cell, last_row_used)."""
    h2(ws, r, header)
    rr = r + 1

    def kv(label, value, fmt, fill, center=False, bold=False):
        nonlocal rr
        a = ws.cell(rr, 1, label); a.border = BORDER
        b = ws.cell(rr, 2, value); b.number_format = fmt; b.fill = fill; b.border = BORDER
        if center:
            b.alignment = Alignment(horizontal="center")
        if bold:
            a.font = BOLD; b.font = BOLD
        cell = "B%d" % rr
        rr += 1
        return cell

    quote_cell = kv(quote_label, quote_formula, FMT_PRICE, CALC_FILL)
    if quote_note:
        note(ws, rr, quote_note, span=3); rr += 1
    coupon_cell = kv("Coupon rate (annual)", coupon, FMT_PCT, INPUT_FILL, center=True)
    kv("Settlement (current) date", settle, FMT_DATE, INPUT_FILL, center=True)
    kv("Last coupon date", last_cpn, FMT_DATE, INPUT_FILL, center=True)
    kv("Next coupon date", next_cpn, FMT_DATE, INPUT_FILL, center=True)
    kv("Day-count convention", "Actual / Actual", "General", INPUT_FILL, center=True)

    # Days since last coupon (accrued-interest numerator) = live sum of monthly day counts
    s = ws.cell(rr, 1, "Days since last coupon (accrued-interest period):"); s.font = BOLD; rr += 1
    first = rr
    for lab, n in days_since_rows:
        kv("    " + lab, n, "0", INPUT_FILL, center=True)
    dsl_cell = kv("    = Days since last coupon", "=SUM(B%d:B%d)" % (first, rr - 1),
                  "0", CALC_FILL, center=True, bold=True)

    # Days in coupon period (denominator) = live sum of monthly day counts
    s = ws.cell(rr, 1, "Days in coupon period:"); s.font = BOLD; rr += 1
    first = rr
    for lab, n in days_period_rows:
        kv("    " + lab, n, "0", INPUT_FILL, center=True)
    dcp_cell = kv("    = Days in coupon period", "=SUM(B%d:B%d)" % (first, rr - 1),
                  "0", CALC_FILL, center=True, bold=True)

    accr_cell = kv("Accrued interest = coupon/2 x 100 x days_since / days_period",
                   "=%s/2*100*%s/%s" % (coupon_cell, dsl_cell, dcp_cell), FMT_ACCR, CALC_FILL)
    full_cell = kv("Full (dirty) price = quoted price + accrued interest",
                   "=%s+%s" % (quote_cell, accr_cell), FMT_PRICE, RESULT_FILL, bold=True)
    return full_cell, rr - 1


tsy_full, tsy_last = price_block(
    s1, 3, "Step 1.1  -  Treasury note 912828EE",
    "Quoted price  97-15+   ( = 97 + 15/32 + 1/64 )", "=97+15/32+1/64",
    "Treasury prices are quoted in 32nds of a point; the trailing '+' adds half a 32nd, i.e. 1/64.",
    TSY_COUPON, SETTLE, TSY_LAST_CPN, TSY_NEXT_CPN, TSY_DAYS_SINCE, TSY_DAYS_PERIOD)

tips_full, tips_last = price_block(
    s1, tsy_last + 3, "Step 1.2  -  TIPS 912828EA",
    "Quoted price  96-20   ( = 96 + 20/32 )", "=96+20/32",
    "Quoted in 32nds (no '+' tick here). This is the real quoted price; the inflation adjustment "
    "is handled separately through the inflation swaps in Step 2.",
    TIPS_COUPON, SETTLE, TIPS_LAST_CPN, TIPS_NEXT_CPN, TIPS_DAYS_SINCE, TIPS_DAYS_PERIOD)

TSY_FULL_REF = "'Step1_Prices'!%s" % tsy_full
TIPS_FULL_REF = "'Step1_Prices'!%s" % tips_full
note(s1, tips_last + 2,
     "Actual/Actual day count: the numerator counts actual days from the last coupon to settlement "
     "(including the last coupon date, excluding the settlement date); the denominator is the actual "
     "number of days in the current coupon period. These full prices are the TIPS cost and the "
     "T-note short proceeds carried into Step2 and Result. Yellow = input, blue = formula, "
     "green = key result.", span=3)

# ===== Sheet 4: Step2_CashFlows =======================================================
s2 = wb.create_sheet("Step2_CashFlows")
title(s2, "Step 2  -  Replicating cash flows: TIPS + inflation swaps + STRIPS", ncols=9)
widths = {"A": 12, "B": 15, "C": 13, "D": 14, "E": 13, "F": 11, "G": 12, "H": 13, "I": 12}
for col, wd in widths.items():
    s2.column_dimensions[col].width = wd

# Parameters / derivation block
h2(s2, 3, "Parameters and derivation (notebook tables 2.1-2.3)")
par = [
    ("TIPS real semi-annual coupon  c_TIPS = 1.875%/2 x 100", "=0.01875/2*100", "B4"),
    ("T-note semi-annual coupon     c_Tnote = 4.250%/2 x 100", "=0.0425/2*100", "B5"),
    ("Face value", 100, "B6"),
]
for lab, val, _a in par:
    rr = int(_a[1:])
    a = s2.cell(rr, 1, lab); a.border = BORDER
    b = s2.cell(rr, 2, val); b.number_format = FMT_PRICE; b.fill = CALC_FILL if isinstance(val, str) else INPUT_FILL
    b.border = BORDER
C_TIPS, C_TNOTE, FACE = "$B$4", "$B$5", "$B$6"
s2.merge_cells("D4:I6")
dcell = s2.cell(4, 4,
    "Per coupon date t:  TIPS + swap = c_TIPS x (1 + P_swap(t)),  with  P_swap(t) = (1+f_t)^t - 1.\n"
    "The floating I_t/I_0 legs of TIPS and swap cancel, so the net is known today.\n"
    "Maturity adds Face to both the (TIPS+swap) and the T-note legs.")
dcell.font = ITAL
dcell.alignment = Alignment(wrap_text=True, vertical="top")

# Initial outflows callout (5 Oct 2006)
h2(s2, 8, "Initial cash flows on the trade date (5 Oct 2006)")
a = s2.cell(9, 1, "Long TIPS"); a.border = BORDER
b = s2.cell(9, 2, "=-%s" % TIPS_FULL_REF); b.number_format = FMT_PRICE; b.fill = CALC_FILL; b.border = BORDER
a = s2.cell(10, 1, "Short T-note"); a.border = BORDER
b = s2.cell(10, 2, "=%s" % TSY_FULL_REF); b.number_format = FMT_PRICE; b.fill = CALC_FILL; b.border = BORDER
s2.cell(9, 3, "(= -TIPS full price)").font = ITAL
s2.cell(10, 3, "(= +T-note full price, received on the short)").font = ITAL

# Live numeric table
TBL_HDR = 13
headers = ["Date", "Swap rate f", "Tenor T (yrs)", "Swap fixed leg\n(1+f)^T-1",
           "TIPS + swap", "T-note", "STRIPS cash flow\n(=T-note-(TIPS+swap))",
           "STRIPS price", "STRIPS cost\n(=cf x price/100)"]
for j, htxt in enumerate(headers, start=1):
    hdr(s2, TBL_HDR, j, htxt)
s2.row_dimensions[TBL_HDR].height = 42

DATA0 = TBL_HDR + 1
for i in range(N):
    r = DATA0 + i
    md_r = MD_DATA0 + i
    last = (i == N - 1)
    # A date
    c = s2.cell(r, 1, "='MarketData'!A%d" % md_r); c.number_format = FMT_DATE; c.border = BORDER
    # B swap rate, C tenor (from MarketData)
    c = s2.cell(r, 2, "='MarketData'!B%d" % md_r); c.number_format = FMT_RATE; c.border = BORDER
    c = s2.cell(r, 3, "='MarketData'!C%d" % md_r); c.number_format = FMT_TENOR; c.border = BORDER
    # D swap fixed leg
    c = s2.cell(r, 4, "=(1+B%d)^C%d-1" % (r, r)); c.number_format = FMT_RATE; c.fill = CALC_FILL; c.border = BORDER
    # E TIPS + swap
    if last:
        f = "=(%s+%s)*(1+D%d)" % (FACE, C_TIPS, r)
    else:
        f = "=%s*(1+D%d)" % (C_TIPS, r)
    c = s2.cell(r, 5, f); c.number_format = FMT_PRICE; c.fill = CALC_FILL; c.border = BORDER
    # F T-note
    f = "=%s+%s" % (FACE, C_TNOTE) if last else "=%s" % C_TNOTE
    c = s2.cell(r, 6, f); c.number_format = FMT_PRICE; c.fill = CALC_FILL; c.border = BORDER
    # G STRIPS cash flow
    c = s2.cell(r, 7, "=F%d-E%d" % (r, r)); c.number_format = FMT_PRICE; c.fill = CALC_FILL; c.border = BORDER
    # H STRIPS price (from MarketData)
    c = s2.cell(r, 8, "='MarketData'!D%d" % md_r); c.number_format = FMT_PRICE; c.border = BORDER
    # I STRIPS cost
    c = s2.cell(r, 9, "=G%d*H%d/100" % (r, r)); c.number_format = FMT_PRICE; c.fill = CALC_FILL; c.border = BORDER
DATA_LAST = DATA0 + N - 1

PX_ROW = DATA_LAST + 1
a = s2.cell(PX_ROW, 8, "PxSTRIPS ="); a.font = BOLD; a.alignment = Alignment(horizontal="right")
b = s2.cell(PX_ROW, 9, "=SUM(I%d:I%d)" % (DATA0, DATA_LAST))
b.number_format = FMT_PRICE; b.font = BOLD; b.fill = RESULT_FILL; b.border = BORDER
PXSTRIPS_REF = "'Step2_CashFlows'!I%d" % PX_ROW
note(s2, PX_ROW + 2,
     "PxSTRIPS is the net cost today of the STRIPS leg (long at coupon dates, short at maturity). "
     "Carried into the Result tab.", span=9)

# ===== Sheet 5: Result ================================================================
rs = wb.create_sheet("Result")
title(rs, "Result  -  Is there an arbitrage?", ncols=4)
rs.column_dimensions["A"].width = 52
rs.column_dimensions["B"].width = 16

def rrow(ws, r, lab, formula, fmt=FMT_PRICE, fill=CALC_FILL, bold=False):
    a = ws.cell(r, 1, lab); a.border = BORDER
    if bold:
        a.font = BOLD
    b = ws.cell(r, 2, formula); b.number_format = fmt; b.fill = fill; b.border = BORDER
    if bold:
        b.font = BOLD
    return b

h2(rs, 3, "Matched-maturity comparison")
rrow(rs, 4, "TIPS market (full) price", "=%s" % TIPS_FULL_REF)
rrow(rs, 5, "plus net cost of STRIPS positions (PxSTRIPS; < 0 = proceeds)", "=%s" % PXSTRIPS_REF)
rrow(rs, 6, "Replicating portfolio cost  =  TIPS + swaps + STRIPS", "=B4+B5", bold=True)
rrow(rs, 8, "Treasury note market (full) price", "=%s" % TSY_FULL_REF)
rrow(rs, 9, "Mispricing  =  Treasury - replicating portfolio", "=B8-B6",
     fill=RESULT_FILL, bold=True)
note(rs, 10,
     "PxSTRIPS is negative: the STRIPS leg is net short (short at maturity) and returns 1.1887 today. "
     "Equivalently the package costs 97.0428 - 1.1887 = 95.8541, exactly as on the lecture slide.", span=2)

h2(rs, 12, "Adjusting for the ~1-month maturity mismatch  (notebook 'Details')")
rrow(rs, 13, "TIPS yield to maturity", TIPS_YIELD, fmt=FMT_PCT, fill=INPUT_FILL)
rrow(rs, 14, "Treasury note repriced at the TIPS yield (4.97%)", TSY_REPRICED_AT_TIPS_YIELD, fill=INPUT_FILL)
rrow(rs, 15, "Mispricing (maturity-adjusted) = Treasury - repriced", "=B8-B14",
     fill=RESULT_FILL, bold=True)
note(rs, 17,
     "TIPS maturity 7/15/2015 vs T-note 8/15/2015. The yield 4.97% and repriced value 95.4583 are "
     "computed in the notebook (Fleckenstein-Longstaff-Lustig 2014); entered here as inputs.", span=2)
note(rs, 19,
     "Arbitrage: short the richer Treasury note at 98.0734 and buy the cheaper replicating package "
     "(TIPS + inflation swaps + STRIPS) at 95.8541, locking in about 2.22 per 100 notional "
     "(about 2.62 after the maturity adjustment).", span=2)

# --------------------------------------------------------------------------------------
# Verification: recompute in pure Python, assert against notebook checkpoints
# --------------------------------------------------------------------------------------
def approx(a, b, tol):
    return abs(a - b) <= tol

c_tips = 0.01875 / 2 * 100      # 0.9375
c_tnote = 0.0425 / 2 * 100      # 2.125
tsy_full_v = (97 + 15/32 + 1/64) + TSY_COUPON/2*100*TSY_DSL/TSY_DCP
tips_full_v = (96 + 20/32) + TIPS_COUPON/2*100*TIPS_DSL/TIPS_DCP

fixedleg = [(1 + CF_SWAPS[i])**CF_TENORS[i] - 1 for i in range(N)]
tips_swap = [(100 + c_tips) * (1 + fixedleg[i]) if i == N-1 else c_tips * (1 + fixedleg[i])
             for i in range(N)]
tnote = [(100 + c_tnote) if i == N-1 else c_tnote for i in range(N)]
strips_cf = [tnote[i] - tips_swap[i] for i in range(N)]
strips_cost = [strips_cf[i] * STRIPS_PX[i] / 100 for i in range(N)]
px_strips = sum(strips_cost)            # signed net cost of STRIPS leg (~ -1.1887; < 0 = net proceeds)
replicating = tips_full_v + px_strips   # TIPS cost + signed STRIPS cost (swaps are costless at inception)
mispricing = tsy_full_v - replicating
mispricing_adj = tsy_full_v - TSY_REPRICED_AT_TIPS_YIELD

# The monthly day-count breakdowns must sum to the accrued numerators/denominators.
assert sum(n for _, n in TSY_DAYS_SINCE) == TSY_DSL, "Treasury days-since breakdown != 51"
assert sum(n for _, n in TSY_DAYS_PERIOD) == TSY_DCP, "Treasury days-period breakdown != 184"
assert sum(n for _, n in TIPS_DAYS_SINCE) == TIPS_DSL, "TIPS days-since breakdown != 82"
assert sum(n for _, n in TIPS_DAYS_PERIOD) == TIPS_DCP, "TIPS days-period breakdown != 184"

checks = [
    ("Treasury full price",          tsy_full_v,     98.07337,  5e-4),
    ("TIPS full price",              tips_full_v,    97.0428,   5e-4),
    ("SwapFixedLeg[2008-01-15]",     fixedleg[2],    0.025183,  5e-5),
    ("PxSTRIPS (signed, net)",       px_strips,     -1.188693,  1e-2),
    ("Replicating cost",             replicating,    95.8541,   1e-2),
    ("Mispricing (matched maturity)",mispricing,     2.22,      1e-2),
    ("Mispricing (maturity-adj)",    mispricing_adj, 2.61507,   5e-4),
]
print("Verification against notebook ground truth")
print("-" * 60)
all_ok = True
for name, got, want, tol in checks:
    ok = approx(got, want, tol)
    all_ok &= ok
    print("  [%s]  %-32s got=%.6f  want=%.6f" % ("PASS" if ok else "FAIL", name, got, want))
print("-" * 60)
if not all_ok:
    print("ERROR: one or more checkpoints failed - workbook NOT saved.")
    sys.exit(1)

# Guard: openpyxl silently turns any cell text starting with "=" into a formula. A label like
# "= Replicating cost ..." then serializes as <f> Replicating cost ...</f>, which Excel rejects and
# removes ("Removed Records: Formula"). Catch any label-misread-as-formula here and fail loudly.
for _ws in wb.worksheets:
    for _row in _ws.iter_rows():
        for _c in _row:
            if _c.data_type == "f" and isinstance(_c.value, str) and _c.value[:2] == "= ":
                raise ValueError("Label misread as formula at %s!%s: %r"
                                 % (_ws.title, _c.coordinate, _c.value))

# Force Excel to recalculate every formula when the file is opened (so cached placeholders are
# replaced with real values immediately), then repair openpyxl's empty <v/> cached-value tags so
# Excel opens the file without a 'corrupt / repaired' prompt.
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
fix_empty_formula_values(OUT)
print("OK: all checkpoints passed.")
print("Saved workbook -> %s" % OUT)
print("Sheets: %s" % ", ".join(wb.sheetnames))
